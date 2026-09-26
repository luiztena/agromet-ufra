"""
comparar_coleta_ano.py
----------------------
Compara a cobertura de coleta (manhã + tarde) entre 2024, 2025 e 2026
até 25/09 de cada ano, para comparação justa.

Uso:
    python scripts/comparar_coleta_ano.py
"""

from openpyxl import load_workbook
from datetime import datetime, time
from pathlib import Path

PLANILHA = Path("dados_brutos/planilhas/Dados TAB.xlsx")
ABAS = [" 2024", "2025", "2026"]
MES_LIMITE = 9
DIA_LIMITE = 25


def analisar(ws):
    por_dia = {}
    for r in ws.iter_rows(values_only=True):
        if not isinstance(r[1], datetime) or not isinstance(r[2], time):
            continue
        # filtra até 25/09
        if r[1].month > MES_LIMITE:
            continue
        if r[1].month == MES_LIMITE and r[1].day > DIA_LIMITE:
            continue

        dia = r[1].strftime("%d/%m")
        por_dia.setdefault(dia, {})[r[2].hour] = r[6]

    completo = sum(1 for h in por_dia.values() if h.get(9) is not None and h.get(15) is not None)
    so_manha = sum(1 for h in por_dia.values() if h.get(9) is not None and h.get(15) is None)
    so_tarde = sum(1 for h in por_dia.values() if h.get(9) is None and h.get(15) is not None)
    ambos_vazios = sum(1 for h in por_dia.values() if h.get(9) is None and h.get(15) is None)

    return {
        "total": len(por_dia),
        "completo": completo,
        "so_manha": so_manha,
        "so_tarde": so_tarde,
        "ambos_vazios": ambos_vazios,
    }


def main():
    wb = load_workbook(PLANILHA, read_only=True, data_only=True)

    print(f"Análise de coleta — período 01/01 a {DIA_LIMITE:02d}/{MES_LIMITE:02d} de cada ano\n")
    print(f"{'Aba':<10} {'Dias':>6} {'Completo':>10} {'Só manhã':>10} {'Só tarde':>10} {'Vazios':>10}")
    print("-" * 60)

    for aba in ABAS:
        if aba not in wb.sheetnames:
            print(f"{aba:<10} (não encontrada)")
            continue
        r = analisar(wb[aba])
        print(f"{aba!r:<10} {r['total']:>6} {r['completo']:>10} {r['so_manha']:>10} {r['so_tarde']:>10} {r['ambos_vazios']:>10}")

    wb.close()


if __name__ == "__main__":
    main()