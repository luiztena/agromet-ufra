"""
regerar_todos_json.py
---------------------
Regenera todos os JSONs a partir da planilha, para cada aba disponível.

- Detecta o nome exato da aba (mesmo que tenha espaço ou caractere estranho).
- Faz backup dos JSONs atuais antes de sobrescrever.
- Roda planilha_para_json.py para cada aba.
- Reporta sucesso/falha por ano.

Uso:
    python scripts\regerar_todos_json.py
"""

import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(r"C:\Users\luizt\Desktop\ISPAAM API\agromet-ufra")
PLANILHA = BASE / "dados_brutos" / "planilhas" / "Dados TAB.xlsx"
DIR_JSON = BASE / "dados_brutos"
DIR_BACKUP = DIR_JSON / "backup_json_antes_saneamento"
SCRIPT_CONVERSOR = BASE / "scripts" / "planilha_para_json.py"

ANOS = list(range(2018, 2027))


def encontrar_aba(abas_reais, ano):
    """
    Procura a aba que corresponde ao ano, ignorando espaços extras.
    Retorna o nome exato como está na planilha, ou None.
    """
    alvo = str(ano)
    for nome in abas_reais:
        if nome.strip() == alvo:
            return nome
    return None


def main():
    print("=" * 70)
    print(" REGERAR TODOS OS JSONs")
    print("=" * 70)

    if not PLANILHA.exists():
        print(f"\nERRO: planilha não encontrada em:\n  {PLANILHA}")
        sys.exit(1)

    if not SCRIPT_CONVERSOR.exists():
        print(f"\nERRO: script conversor não encontrado em:\n  {SCRIPT_CONVERSOR}")
        sys.exit(1)

    # Lê as abas
    wb = load_workbook(PLANILHA, read_only=True)
    abas_reais = wb.sheetnames
    wb.close()

    print(f"\n Planilha: {PLANILHA.name}")
    print(f" Abas encontradas ({len(abas_reais)}):")
    for nome in abas_reais:
        # Mostra com marcação visível para espaços
        marcado = nome.replace(" ", "·")
        print(f"   - '{nome}'  (repr: {repr(nome)}, marcado: {marcado})")

    print(f"\n Backup para: {DIR_BACKUP}")

    # Mapeia cada ano para o nome real da aba
    mapa_anos = {}
    for ano in ANOS:
        nome_real = encontrar_aba(abas_reais, ano)
        if nome_real:
            mapa_anos[ano] = nome_real
        else:
            print(f"\n [!] Ano {ano} não tem aba correspondente — será pulado")

    if not mapa_anos:
        print("\nNenhuma aba válida encontrada. Abortando.")
        return

    print("\n Mapeamento ano -> aba:")
    for ano, nome in mapa_anos.items():
        diff = " (nome exato)" if nome == str(ano) else f"  <<< '{nome}'"
        print(f"   {ano} -> {diff.strip()}")

    resposta = input("\n Prosseguir com a regeneração? (s/N): ").strip().lower()
    if resposta != "s":
        print(" Cancelado.")
        return

    # Backup
    DIR_BACKUP.mkdir(parents=True, exist_ok=True)
    n_backups = 0
    for ano in mapa_anos:
        origem = DIR_JSON / f"dados_{ano}.json"
        if origem.exists():
            destino = DIR_BACKUP / f"dados_{ano}.json"
            shutil.copy2(origem, destino)
            n_backups += 1
    print(f"\n Backup: {n_backups} arquivos copiados para\n   {DIR_BACKUP}")

    # Regerar cada ano
    print("\n Regerando JSONs:")
    sucessos = 0
    falhas = 0

    for ano, nome_aba in sorted(mapa_anos.items()):
        destino = DIR_JSON / f"dados_{ano}.json"

        cmd = [
            sys.executable,
            str(SCRIPT_CONVERSOR),
            str(PLANILHA),
            "--aba", nome_aba,
            "-o", str(destino),
        ]

        print(f"\n  [{ano}] aba '{nome_aba}'")

        # encoding="cp1252"  -> Windows padrão
        # errors="replace"   -> não quebra se aparecer byte estranho
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="cp1252",
            errors="replace",
        )

        if result.returncode != 0:
            print(f"  [{ano}] FALHA (returncode={result.returncode})")
            stderr = (result.stderr or "").strip()
            stdout = (result.stdout or "").strip()
            if stderr:
                print(f"    stderr: {stderr[:400]}")
            if stdout:
                print(f"    stdout: {stdout[:400]}")
            falhas += 1
            continue

        # Extrai as últimas linhas úteis do stdout
        stdout = result.stdout or ""
        linhas = [l for l in stdout.strip().splitlines() if l.strip()]
        for l in linhas[-3:]:
            print(f"    {l}")
        sucessos += 1

    # Síntese
    print("\n" + "=" * 70)
    print(" FIM")
    print("=" * 70)
    print(f"  Sucessos: {sucessos}")
    print(f"  Falhas:   {falhas}")
    print(f"  Backup:   {DIR_BACKUP}")


if __name__ == "__main__":
    main()