"""
debug_contagem.py
-----------------
Simula o processamento completo do planilha_para_json.py para uma aba,
mostrando quantas linhas são processadas, quantas ignoradas, e quais
dias aparecem no final.
"""

import sys
from openpyxl import load_workbook

sys.path.insert(0, "scripts")
from planilha_para_json import formatar_valor, limpar_cabecalho

PLANILHA = r"dados_brutos\planilhas\Dados TAB.xlsx"
ABA = " 2024"

wb = load_workbook(PLANILHA, read_only=True, data_only=True)
ws = wb[ABA]

linhas_iter = ws.iter_rows(values_only=True)
cab_raw = next(linhas_iter)
cab_limpo = [limpar_cabecalho(c, i) for i, c in enumerate(cab_raw)]

print(f"Cabeçalho: {len(cab_limpo)} colunas")

registros = []
ignorados = 0
processadas = 0

for linha in linhas_iter:
    processadas += 1
    if all(c is None or (isinstance(c, str) and c.strip() == "") for c in linha):
        ignorados += 1
        continue

    obj = {}
    for i, valor in enumerate(linha):
        if i >= len(cab_limpo):
            break
        obj[cab_limpo[i]] = formatar_valor(valor)

    if not obj.get("Data    (dd/mm/aaaa)"):
        ignorados += 1
        continue

    registros.append(obj)

print(f"Processadas: {processadas}")
print(f"Ignoradas: {ignorados}")
print(f"Registros finais: {len(registros)}")

print()
print("Datas 13/03 encontradas:")
for r in registros:
    if r.get("Data    (dd/mm/aaaa)") == "13/03":
        print(f"  {r.get('Data    (dd/mm/aaaa)')} {r.get('Hora local (hh:mm)')} Tar={r.get('Tar (°C)')}")