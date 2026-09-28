"""
corrigir_datas.py
-----------------
Corrige datas erradas nas abas da planilha, contornando o problema
de locale do Excel (que interpreta datas ambíguas como mm/dd).

Edita diretamente via openpyxl. Grava valores exatos sem ambiguidade.

Uso:
    python scripts/corrigir_datas.py
"""

from datetime import datetime
from openpyxl import load_workbook

PLANILHA = r"dados_brutos\planilhas\Dados TAB.xlsx"

# Estrutura: {aba: [(linha, ano, mes, dia, descricao), ...]}
CORRECOES = {
    "2019": [
        (51, 2019, 1, 25, "25/01/2019 (texto) -> 25/01/2019 (data)"),
    ],
    "2020": [
        (678, 2020, 4, 12, "04/12/2020 -> 12/04/2020 (dia/mês invertidos)"),
    ],

    # Descomente e adicione outras abas conforme necessário:
    # "2021": [...],
    # "2022": [...],
    # "2023": [...],
}


def main():
    print(f"Abrindo: {PLANILHA}")
    wb = load_workbook(PLANILHA)

    total_correcoes = 0

    for aba, correcoes in CORRECOES.items():
        if aba not in wb.sheetnames:
            print(f"  ⚠ Aba {aba!r} não encontrada. Pulando.")
            continue

        ws = wb[aba]
        print(f"\nAba {aba!r}: {len(correcoes)} correções")

        for linha, ano, mes, dia, descricao in correcoes:
            celula = ws.cell(row=linha, column=2)  # coluna 2 = B (Data)
            valor_antigo = celula.value
            celula.value = datetime(ano, mes, dia)
            print(f"  Linha {linha}: {valor_antigo!r} -> {celula.value!r}")
            print(f"    ({descricao})")
            total_correcoes += 1

    wb.save(PLANILHA)
    wb.close()
    print(f"\n✅ Planilha salva. {total_correcoes} correções aplicadas.")


if __name__ == "__main__":
    main()