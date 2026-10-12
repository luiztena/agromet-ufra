"""
planilha_para_json.py
---------------------
Converte uma aba específica de uma planilha .xlsx no JSON bruto do ISPAAM.

Uso:
    python scripts/planilha_para_json.py "dados_brutos/planilhas/Dados TAB.xlsx" --aba 2026 -o dados_brutos/dados_2026.json
"""

import argparse
import json
import sys
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Mapeamentos de cabeçalho
# ---------------------------------------------------------------------------

ALIASES_CABECALHO = {
    "Data (dd/mm/aaaa)": "Data    (dd/mm/aaaa)",
    "Hora (hh:mm)": "Hora local (hh:mm)",
    "Prp (mm)": "Prp (mm/dia)",
    "Ev (mm) coleta": "Ev (mm)",
    "Ev (mm)*      coleta": "Ev (mm)",
    "Ev (mm)*        após enchimento": "Ev (mm)*",
    "Ev (mm)": "Ev (mm/dia)",
    "Observações": "Observadores",
}

MAPA_2018 = {
    0: "Dia da semana",
    1: "Data    (dd/mm/aaaa)",
    2: "Hora local (hh:mm)",
    3: "Hora UTC (hh:mm)",
    4: "Tmáx (°C)",
    5: "Tmin (°C)",
    6: "Tméd (ºC)",
    7: "Tar (°C)",
    8: "TH2O ev (°C)",
    9: "U2 (m/s)",
    10: "Tar 2 (°C)",
    11: "Prp (mm/dia)",
    12: "Ev (mm)",
    13: "Ev (mm)*",
    14: "Patm (mbar)",
    15: "Tmáx Real (ºC)",
    16: "Tmin Real (ºC)",
    17: "esTU",
    18: "ea",
    19: "es",
    20: "UR (%)",
    21: "Ev (mm/dia)",
    22: "Observadores",
}


# ---------------------------------------------------------------------------
# Normalizacao de dia da semana (para validacao)
# ---------------------------------------------------------------------------

_DIAS_SEMANA = ["segunda", "terca", "quarta", "quinta", "sexta", "sabado", "domingo"]

_MAPA_DIA_SEMANA = {
    "seg": "segunda", "segunda": "segunda", "segunda-feira": "segunda",
    "ter": "terca",   "terca": "terca",     "terca-feira": "terca",
    "qua": "quarta",  "quarta": "quarta",   "quarta-feira": "quarta",
    "qui": "quinta",  "quinta": "quinta",   "quinta-feira": "quinta",
    "sex": "sexta",   "sexta": "sexta",     "sexta-feira": "sexta",
    "sab": "sabado",  "sabado": "sabado",
    "dom": "domingo", "domingo": "domingo",
}


def _normalizar_dia_semana(s):
    """Normaliza 'seg' -> 'segunda', 'Sexta-feira' -> 'sexta', etc."""
    if not s:
        return None
    s = s.lower().strip()
    s = (s.replace("á", "a").replace("ã", "a").replace("â", "a")
           .replace("é", "e").replace("ê", "e")
           .replace("í", "i")
           .replace("ó", "o").replace("ô", "o").replace("õ", "o")
           .replace("ú", "u")
           .replace("ç", "c"))
    return _MAPA_DIA_SEMANA.get(s, s)


def _dia_semana_confere(data_dd_mm, dia_semana_str, ano):
    """
    Verifica se o dia da semana bate com a data (formato dd/mm).
    Retorna True se OK, False se inconsistente, True se nao der pra validar.
    """
    if not data_dd_mm or not dia_semana_str or not ano:
        return True

    partes = str(data_dd_mm).split("/")
    if len(partes) < 2:
        return True

    try:
        dd = int(partes[0])
        mm = int(partes[1])
        data = datetime(ano, mm, dd)
    except (ValueError, IndexError, TypeError):
        return True  # data invalida, deixa passar

    dia_real = _DIAS_SEMANA[data.weekday()]
    dia_planilha = _normalizar_dia_semana(dia_semana_str)

    if dia_planilha is None:
        return True

    return dia_planilha == dia_real


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def formatar_valor(valor):
    """Converte um valor de célula do Excel no formato que queremos no JSON."""
    if valor is None:
        return None

    if isinstance(valor, datetime):
        if valor.hour == 0 and valor.minute == 0 and valor.second == 0:
            return f"{valor.day:02d}/{valor.month:02d}"
        else:
            return valor.strftime("%d/%m %H:%M")

    if isinstance(valor, time):
        return f"{valor.hour:02d}:{valor.minute:02d}"

    if isinstance(valor, str):
        s = valor.strip()
        return s if s else None

    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and valor.is_integer():
            return str(int(valor))
        return str(valor).replace(".", ",")

    return str(valor)


# ---------------------------------------------------------------------------
# Saneamento pós-conversão
# ---------------------------------------------------------------------------

def _para_float_local(v):
    """Converte string com vírgula ou ponto para float. None se não numérico."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s.lower() == "null" or s in ("-", "--"):
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def _formatar_decimal(v):
    """Float -> string com vírgula decimal (padrão dos JSONs)."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(round(v, 2)).replace(".", ",")


def sanear_registro(obj):
    """
    Recalcula Tméd e anula valores implausíveis.
    Corrige o bug do Excel: (vazio + número) / 2 gera valores absurdos (16-18 °C).

    Regras:
      - Tméd só é calculada se Tmin e Tmáx existem (prioridade Real).
      - Tmin fora de [15, 27] °C vira None.
      - Tmáx fora de [25, 40] °C vira None.
      - Tméd fora de [18, 40] °C vira None (redundância de segurança).
    """
    # Tmin e Tmáx com prioridade para os "Real"
    tmin = _para_float_local(obj.get("Tmin Real (ºC)"))
    if tmin is None:
        tmin = _para_float_local(obj.get("Tmin (°C)"))

    tmax = _para_float_local(obj.get("Tmáx Real (ºC)"))
    if tmax is None:
        tmax = _para_float_local(obj.get("Tmáx (°C)"))

    # Faixas plausíveis para Belém
    if tmin is not None and not (15.0 <= tmin <= 27.0):
        tmin = None
    if tmax is not None and not (25.0 <= tmax <= 40.0):
        tmax = None

    # Recalcula Tméd
    if tmin is not None and tmax is not None:
        tmed = (tmin + tmax) / 2
        if not (18.0 <= tmed <= 40.0):
            tmed = None
    else:
        tmed = None

    # Aplica ao objeto
    if "Tmin Real (ºC)" in obj:
        obj["Tmin Real (ºC)"] = _formatar_decimal(tmin)
    if "Tmin (°C)" in obj:
        obj["Tmin (°C)"] = _formatar_decimal(tmin)
    if "Tmáx Real (ºC)" in obj:
        obj["Tmáx Real (ºC)"] = _formatar_decimal(tmax)
    if "Tmáx (°C)" in obj:
        obj["Tmáx (°C)"] = _formatar_decimal(tmax)
    if "Tméd (ºC)" in obj:
        obj["Tméd (ºC)"] = _formatar_decimal(tmed)

    return obj


# ---------------------------------------------------------------------------
# Cabeçalho
# ---------------------------------------------------------------------------

def normalizar_cabecalho(nome, idx):
    if nome is None:
        return f"coluna_{idx}"
    if not isinstance(nome, str):
        return str(nome)

    nome_limpo = " ".join(nome.split()) if nome.strip() else nome

    if nome in ALIASES_CABECALHO:
        return ALIASES_CABECALHO[nome]
    if nome_limpo in ALIASES_CABECALHO:
        return ALIASES_CABECALHO[nome_limpo]

    return nome


def obter_cabecalho(cabecalho_raw, aba):
    if aba and aba.strip() == "2018":
        cabecalho = []
        for i in range(len(cabecalho_raw)):
            cabecalho.append(MAPA_2018.get(i, f"coluna_{i}"))
        print(f"  [2018] Usando mapeamento posicional especial")
        return cabecalho

    return [normalizar_cabecalho(c, i) for i, c in enumerate(cabecalho_raw)]


def coluna_data_canonica(aba):
    return "Data    (dd/mm/aaaa)"


# ---------------------------------------------------------------------------
# Conversão
# ---------------------------------------------------------------------------

def converter(caminho_entrada: Path, caminho_saida: Path, aba: str = None):
    if not caminho_entrada.exists():
        print(f"Arquivo não encontrado: {caminho_entrada}")
        sys.exit(1)

    # Extrai o ano da aba para validar dia da semana
    ano_aba = None
    if aba:
        try:
            ano_aba = int(aba.strip())
        except ValueError:
            ano_aba = None

    print(f"Lendo: {caminho_entrada}")
    wb = load_workbook(caminho_entrada, data_only=True, read_only=True)

    if aba:
        if aba not in wb.sheetnames:
            print(f"Aba '{aba}' não existe. Disponíveis: {wb.sheetnames}")
            sys.exit(1)
        ws = wb[aba]
    else:
        ws = wb[wb.sheetnames[0]]
        print(f"Usando primeira aba: '{ws.title}'")

    print(f"Aba usada: '{ws.title}'")
    if ano_aba:
        print(f"Ano inferido (para validação): {ano_aba}")

    linhas_iter = ws.iter_rows(values_only=True)

    try:
        cabecalho_raw = next(linhas_iter)
    except StopIteration:
        print("Planilha vazia.")
        sys.exit(1)

    cabecalho = obter_cabecalho(cabecalho_raw, ws.title)
    print(f"Colunas ({len(cabecalho)}): {cabecalho[:6]}...")

    coluna_data = coluna_data_canonica(ws.title)

    registros = []
    ignorados = 0
    rejeitados_dia_semana = 0

    for linha in linhas_iter:
        if all(c is None or (isinstance(c, str) and c.strip() == "") for c in linha):
            ignorados += 1
            continue

        obj = {}
        for i, valor in enumerate(linha):
            if i >= len(cabecalho):
                break
            obj[cabecalho[i]] = formatar_valor(valor)

        if not obj.get(coluna_data):
            ignorados += 1
            continue

        # Validação: dia da semana bate com a data?
        data_str = obj.get(coluna_data)
        dia_semana_str = obj.get("Dia da semana")
        if not _dia_semana_confere(data_str, dia_semana_str, ano_aba):
            print(f"  [SKIP] {data_str} ({dia_semana_str}) — dia da semana inconsistente")
            rejeitados_dia_semana += 1
            continue

        # Saneamento
        obj = sanear_registro(obj)

        registros.append(obj)

    wb.close()

    with open(caminho_saida, "w", encoding="utf-8") as f:
        json.dump(registros, f, ensure_ascii=False, indent=4)

    print(f"\n  Registros gerados: {len(registros)}")
    if ignorados:
        print(f"  Linhas vazias ignoradas: {ignorados}")
    if rejeitados_dia_semana:
        print(f"  Linhas rejeitadas (dia da semana inconsistente): {rejeitados_dia_semana}")
    print(f"  Salvo em: {caminho_saida}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Converte uma aba de planilha .xlsx em JSON bruto."
    )
    parser.add_argument("entrada", help="Caminho da planilha .xlsx.")
    parser.add_argument("-o", "--saida",
                        help="JSON de saída. Padrão: dados_brutos/dados_<aba>.json")
    parser.add_argument("--aba", help="Nome da aba (padrão: primeira).")
    args = parser.parse_args()

    entrada = Path(args.entrada)

    if args.saida:
        saida = Path(args.saida)
    elif args.aba:
        saida = Path("dados_brutos") / f"dados_{args.aba.strip()}.json"
    else:
        saida = Path("dados_brutos") / f"dados_{entrada.stem}.json"

    saida.parent.mkdir(parents=True, exist_ok=True)
    converter(entrada, saida, aba=args.aba)


if __name__ == "__main__":
    main()