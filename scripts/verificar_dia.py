"""
verificar_dia.py
----------------
Verifica quantas observações existem para um dia específico,
tanto no JSON quanto na planilha. Útil pra investigar dias incompletos.

Uso:
    python scripts/verificar_dia.py dados_brutos/dados_2024.json " 2024" 13/03
"""

import json
import sys
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook


def verificar_json(caminho_json, data_alvo):
    print(f"\n=== JSON: {caminho_json} ===")
    with open(caminho_json, encoding="utf-8") as f:
        dados = json.load(f)

    encontrados = [
        x for x in dados
        if x.get("Data    (dd/mm/aaaa)") == data_alvo
    ]

    print(f"Linhas com Data = {data_alvo!r}: {len(encontrados)}")
    for x in encontrados:
        print(f"  Hora = {x.get('Hora local (hh:mm)')!r} | Tar = {x.get('Tar (°C)')!r} | Obs = {x.get('Observadores')!r}")


def verificar_planilha(caminho_planilha, nome_aba, data_alvo):
    print(f"\n=== Planilha: {caminho_planilha} (aba {nome_aba!r}) ===")
    wb = load_workbook(caminho_planilha, read_only=True, data_only=True)

    if nome_aba not in wb.sheetnames:
        print(f"Aba {nome_aba!r} não encontrada. Disponíveis: {wb.sheetnames}")
        return

    ws = wb[nome_aba]
    dia, mes = data_alvo.split("/")

    linhas = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if len(r) < 3:
            continue
        data_cel = r[1]
        if isinstance(data_cel, datetime) and data_cel.month == int(mes) and data_cel.day == int(dia):
            tar = r[6] if len(r) > 6 else None
            obs = r[30] if len(r) > 30 else None
            hora = r[2] if len(r) > 2 else None
            linhas.append((i, data_cel, hora, tar, obs))

    print(f"Linhas com {data_alvo} na aba: {len(linhas)}")
    for i, d, h, tar, obs in linhas:
        print(f"  Linha {i}: {d.strftime('%d/%m/%Y')} {h} | Tar = {tar!r} | Obs = {obs!r}")


def main():
    if len(sys.argv) < 4:
        print("Uso: python scripts/verificar_dia.py <json> <aba> <dd/mm>")
        print("Ex:  python scripts/verificar_dia.py dados_brutos/dados_2024.json ' 2024' 13/03")
        sys.exit(1)

    caminho_json = Path(sys.argv[1])
    nome_aba = sys.argv[2]
    data_alvo = sys.argv[3]

    if caminho_json.exists():
        verificar_json(caminho_json, data_alvo)
    else:
        print(f"JSON não encontrado: {caminho_json}")

    planilha = Path("dados_brutos/planilhas/Dados TAB.xlsx")
    if planilha.exists():
        verificar_planilha(planilha, nome_aba, data_alvo)
    else:
        print(f"Planilha não encontrada: {planilha}")


if __name__ == "__main__":
    main()