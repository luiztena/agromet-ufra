"""
debug_linha.py
--------------
Debug: mostra o que acontece com uma linha específica ao passar
pelo mesmo processamento do planilha_para_json.py.

Uso:
    python scripts/debug_linha.py "dados_brutos/planilhas/Dados TAB.xlsx" " 2024" 146
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

# Importa as funções do script original
sys.path.insert(0, "scripts")
from planilha_para_json import formatar_valor, limpar_cabecalho


def main():
    if len(sys.argv) < 4:
        print("Uso: python scripts/debug_linha.py <planilha> <aba> <numero_linha>")
        sys.exit(1)

    caminho = Path(sys.argv[1])
    nome_aba = sys.argv[2]
    num_linha = int(sys.argv[3])

    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[nome_aba]

    linhas_iter = ws.iter_rows(values_only=True)
    cab_raw = next(linhas_iter)
    cab_limpo = [limpar_cabecalho(c, i) for i, c in enumerate(cab_raw)]

    print(f"Cabeçalho ({len(cab_limpo)} colunas):")
    for i, c in enumerate(cab_limpo):
        print(f"  {i}: {c!r}")

    print()
    print(f"Procurando linha {num_linha}...")

    i = 1
    encontrou = False
    for i, linha in enumerate(linhas_iter, start=2):
        if i == num_linha:
            encontrou = True
            print(f"Linha {i} encontrada:")
            print(f"  Comprimento: {len(linha)}")

            obj = {}
            for j, valor in enumerate(linha):
                if j >= len(cab_limpo):
                    break
                obj[cab_limpo[j]] = formatar_valor(valor)

            print(f"  Chaves no obj: {len(obj)}")
            print(f"  Tem 'Data    (dd/mm/aaaa)'? {'Data    (dd/mm/aaaa)' in obj}")
            print(f"  Valor da Data: {obj.get('Data    (dd/mm/aaaa)')!r}")
            print(f"  Data é truthy? {bool(obj.get('Data    (dd/mm/aaaa)'))}")

            print()
            print("  Campos preenchidos:")
            for k, v in obj.items():
                if v is not None:
                    print(f"    {k!r}: {v!r}")

            break

    if not encontrou:
        print(f"Linha {num_linha} não encontrada (só {i} linhas lidas).")

    wb.close()


if __name__ == "__main__":
    main()