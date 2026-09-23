"""
corrigir_ano_planilha.py
------------------------
Corrige o ano das datas de uma aba específica de uma planilha .xlsx.

Problema que resolve:
    Uma aba chamada '2026' tem datas internas marcadas como 2020
    (por exemplo, por cópia de dados antigos ou formato do Excel).
    O ano interno do datetime não bate com o nome da aba.

O que faz:
    - Lê uma cópia da planilha (o original NUNCA é alterado).
    - Para cada linha da aba alvo, se a coluna Data contém um datetime
      com ano diferente do esperado, substitui por:
          datetime(ano_esperado, mes, dia)
      mantendo mês e dia originais.
    - Salva em um novo arquivo (`Dados TAB - corrigida.xlsx`).

Uso:
    python scripts/corrigir_ano_planilha.py "dados_brutos\planilhas\Dados TAB.xlsx" --aba 2026 --ano 2026

    --aba   Nome da aba a corrigir (ex.: 2026)
    --ano   Ano esperado nas datas (default: extraído do nome da aba)
    -o      Arquivo de saída (default: adiciona ' - corrigida' ao nome)
"""

import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def extrair_ano_do_nome(nome: str):
    """Tenta extrair um ano (4 dígitos) do nome da aba."""
    m = re.search(r"\b(\d{4})\b", nome)
    return int(m.group(1)) if m else None


def corrigir_aba(ws, ano_esperado: int, coluna_data_idx: int = 1):
    """
    Corrige as datas da coluna `coluna_data_idx` (0-based) na aba.

    Retorna (n_corrigidas, n_ja_ok, n_outros).
    """
    n_corrigidas = 0
    n_ja_ok = 0
    n_outros = 0

    # Iteramos a partir da linha 2 (linha 1 é cabeçalho)
    for row in ws.iter_rows(min_row=2):
        if coluna_data_idx >= len(row):
            continue
        celula = row[coluna_data_idx]
        valor = celula.value

        if valor is None:
            continue

        if isinstance(valor, datetime):
            if valor.year == ano_esperado:
                n_ja_ok += 1
            else:
                # Substitui mantendo mês, dia (e hora, se houver)
                celula.value = datetime(
                    ano_esperado,
                    valor.month,
                    valor.day,
                    valor.hour,
                    valor.minute,
                    valor.second,
                )
                n_corrigidas += 1
        else:
            # Valor que não é datetime (string, número, etc.)
            n_outros += 1

    return n_corrigidas, n_ja_ok, n_outros


def corrigir_planilha(caminho_entrada: Path, caminho_saida: Path,
                      aba: str, ano_esperado: int):
    if not caminho_entrada.exists():
        print(f"Arquivo não encontrado: {caminho_entrada}")
        sys.exit(1)

    print(f"Lendo: {caminho_entrada}")
    print(f"Aba alvo: {aba}")
    print(f"Ano esperado: {ano_esperado}")
    print()

    # Faz uma cópia primeiro (garantia extra)
    print(f"Copiando para: {caminho_saida}")
    shutil.copy2(caminho_entrada, caminho_saida)

    # Abre a cópia e edita
    wb = load_workbook(caminho_saida)

    if aba not in wb.sheetnames:
        print(f"Aba '{aba}' não existe. Disponíveis: {wb.sheetnames}")
        wb.close()
        sys.exit(1)

    ws = wb[aba]

    print(f"Corrigindo datas na coluna B (índice 1)...")
    n_corrigidas, n_ja_ok, n_outros = corrigir_aba(ws, ano_esperado)

    wb.save(caminho_saida)
    wb.close()

    print()
    print(f"  Datas corrigidas:  {n_corrigidas}")
    print(f"  Datas já corretas: {n_ja_ok}")
    print(f"  Outros valores:    {n_outros} (não eram datas)")
    print()
    print(f"Salvo em: {caminho_saida}")
    print()
    print("IMPORTANTE: o arquivo original NÃO foi alterado.")
    print("Confira o arquivo corrigido no Excel antes de substituir.")


def main():
    parser = argparse.ArgumentParser(
        description="Corrige o ano das datas em uma aba de planilha .xlsx."
    )
    parser.add_argument("entrada", help="Caminho da planilha .xlsx original.")
    parser.add_argument("--aba", required=True, help="Nome da aba a corrigir (ex.: 2026).")
    parser.add_argument("--ano", type=int,
                        help="Ano esperado. Default: extraído do nome da aba.")
    parser.add_argument("-o", "--saida",
                        help="Caminho do arquivo de saída. Default: '<entrada> - corrigida.xlsx'.")
    args = parser.parse_args()

    entrada = Path(args.entrada)
    if not entrada.exists():
        print(f"Arquivo não encontrado: {entrada}")
        sys.exit(1)

    ano_esperado = args.ano
    if ano_esperado is None:
        ano_esperado = extrair_ano_do_nome(args.aba)
        if ano_esperado is None:
            print(f"Não foi possível inferir o ano a partir da aba '{args.aba}'. Use --ano.")
            sys.exit(1)

    if args.saida:
        saida = Path(args.saida)
    else:
        saida = entrada.with_name(f"{entrada.stem} - corrigida{entrada.suffix}")

    corrigir_planilha(entrada, saida, args.aba, ano_esperado)


if __name__ == "__main__":
    main()